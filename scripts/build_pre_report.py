"""Build the explanatory DOCX with the bundled document runtime, after experiments."""

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]


def build():
    summary = json.loads((ROOT / "results/benchmark/summary.json").read_text())
    coverage = json.loads((ROOT / "results/test-coverage.json").read_text())
    with (ROOT / "results/sensitivity.csv").open() as handle:
        sensitivity = list(csv.DictReader(handle))
    doc = Document()
    # Some bundled Word defaults add a border to the Title style.
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    for name in ("Normal", "List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.06
    for name, size in (("Title", 24), ("Heading 1", 17), ("Heading 2", 12.5)):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(6)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("Project pre-report  |  ").font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "5G UE Context Desynchronization Detection and Recovery"
    doc.core_properties.subject = "Project pre-report and viva preparation"
    doc.core_properties.author = "ACN Project Team"

    def p(text, style=None):
        return doc.add_paragraph(text, style)

    def h(text):
        doc.add_heading(text, level=2)

    def page(title):
        doc.add_page_break()
        doc.add_heading(title, level=1)

    def bullets(items):
        for item in items:
            p(item, "List Bullet")

    def table(headers, rows, widths):
        t = doc.add_table(rows=1, cols=len(headers))
        t.autofit = False
        for i, width in enumerate(widths):
            t.columns[i].width = Inches(width)
        for i, text in enumerate(headers):
            t.rows[0].cells[i].text = text
        for row in rows:
            cells = t.add_row().cells
            for i, text in enumerate(row):
                cells[i].text = str(text)
        for index, row in enumerate(t.rows):
            for i, cell in enumerate(row.cells):
                cell.width = Inches(widths[i])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                properties = cell._tc.get_or_add_tcPr()
                shading = OxmlElement("w:shd")
                shading.set(
                    qn("w:fill"),
                    "203B55" if index == 0 else ("F2F5F8" if index % 2 == 0 else "FFFFFF"),
                )
                properties.append(shading)
                borders = OxmlElement("w:tcBorders")
                for edge in ("top", "left", "bottom", "right"):
                    element = OxmlElement(f"w:{edge}")
                    for attr, val in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                        element.set(qn(f"w:{attr}"), val)
                    borders.append(element)
                properties.append(borders)
                margin = OxmlElement("w:tcMar")
                for edge in ("top", "bottom", "left", "right"):
                    element = OxmlElement(f"w:{edge}")
                    element.set(qn("w:w"), "90")
                    element.set(qn("w:type"), "dxa")
                    margin.append(element)
                properties.append(margin)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.space_before = Pt(2)
                    for run in paragraph.runs:
                        run.font.size = Pt(10)
                        if index == 0:
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
            # Prevent a single row being split across pages.
            cant = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(cant)
        repeated = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(repeated)
        p("")

    def ref_link(paragraph, label, url):
        from docx.opc.constants import RELATIONSHIP_TYPE

        rel_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), rel_id)
        run = OxmlElement("w:r")
        props = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "215A84")
        props.append(color)
        run.append(props)
        text = OxmlElement("w:t")
        text.text = label
        run.append(text)
        link.append(run)
        paragraph._p.append(link)

    def select(scenario, recovery=False, mode="aware"):
        return next(
            r
            for r in summary
            if r["scenario"] == scenario and r["recovery"] == recovery and r["mode"] == mode
        )

    p("5G UE Context Desynchronization Detection and Recovery", "Title")
    p("Project pre-report and viva preparation\n7 October 2026")
    h("Project in one paragraph")
    p(
        "This project implements a protocol-aware verifier for an abstract 5G network. It tracks the RAN's applied resources and the Core's committed session information for each UE/session, waits for valid procedures to finish, and detects persistent contradictions. It then evaluates a bounded reconciliation policy using controlled faults. The completed deliverable is a reproducible simulation prototype with a dashboard, tests, traces and measured comparisons."
    )
    h("The main question")
    p(
        "Can a monitor distinguish normal temporary RAN–Core differences from persistent context disagreement, and can reconciliation reduce the time that disagreement remains? The experiment answers this within an explicit state model. It does not establish performance in an operator network."
    )
    h("Scope and meaning")
    bullets(
        [
            "Model session establishment, QoS resource modification and release. Model access deactivation/resumption and a local cell change as controls.",
            "Use separate Core and RAN state objects, complete telemetry snapshots, procedure deadlines and transaction identities.",
            "Detect three contradiction types and evaluate simulated recovery. No real UE, gNB, RF channel, NGAP encoding or 5G Core is deployed.",
            "Recovery success means context convergence. Removing orphan RAN resources can restore consistency while leaving a user without an established session.",
        ]
    )
    table(
        ["Term", "Meaning in this project"],
        [
            ("UE", "A user device, with one or more modeled sessions"),
            ("RAN and gNB", "Radio access network and its base station"),
            ("Core", "An abstraction combining session and mobility responsibilities"),
            (
                "PDU session",
                "A logical data session; it can persist without active access resources",
            ),
            ("QoS profile", "An abstract integer configuration identifier, not a QFI or 5QI"),
            ("FSM", "A finite-state model whose events change stored context"),
        ],
        [1.25, 5.65],
    )

    page("Research basis and project contributions")
    h("Conceptual base reference")
    p(
        "Data Consistency in the 5G Specification, ICIN 2020, is the conceptual reference selected for this implementation [1]. It examines how 5G procedures access and update UE context, failure scenarios affecting correctness, and consistency properties needed from the data store. It provides a basis for treating context agreement as a distributed-systems correctness problem."
    )
    p(
        "Our extension turns that motivation into an executable monitor and evaluation harness for a narrower RAN–Core model. The accessible abstract and indexed introduction support this positioning. Full-paper retrieval was unsuccessful, so we do not present a complete paper review, reproduce its experiments, or attribute an unverified limitation statement to its authors."
    )
    h("Supporting anomaly detection papers")
    p(
        "The CSCN 2023 paper uses an LSTM autoencoder on CPU and memory metrics during UE registration to identify stress-related anomalies [2]. It motivates anomaly monitoring, but our input is per-session state rather than resource utilization. This is a difference in task and evidence, not proof that its model cannot identify related failures."
    )
    p(
        "The ICTC 2025 paper addresses early detection on large-scale Core KPI streams using point-level amplification and automatic score thresholds [3]. Early detection alone is therefore not our novelty. Our distinguishing engineering goal is identifying a specific context contradiction with protocol evidence and evaluating reconciliation. These comparisons use publisher summaries rather than full-text reproduction."
    )
    h("Contributions we can defend")
    bullets(
        [
            "Per-session semantic rules compare shared meaning across separate state models, including retained sessions with deactivated access resources.",
            "Procedure-aware checking combines pending operations, freshness, sequence ordering and overlapping observation intervals before requiring persistence.",
            "Each alarm records the violated rule, session identity, both observed contexts and timing evidence.",
            "A reproducible harness compares two simpler monitors and measures bounded recovery, including outage and adapter-failure cases.",
        ]
    )
    p(
        "These are implemented contributions relative to the selected references. We do not claim the first state verifier, a new 3GPP procedure, or a proven worldwide research novelty."
    )

    page("State model and detection logic")
    h("What each side stores")
    p(
        "The RAN stores applied access resources, a local retained context, an abstract QoS profile, RRC state and a local cell. The Core stores logical session existence, expected access resources and committed QoS. Both carry a transaction identity, version and change timestamp. A pending Core operation includes a deadline. AMF and SMF responsibilities are aggregated rather than implemented as separate network functions."
    )
    table(
        ["Rule", "Contradiction"],
        [
            ("Orphan RAN resource", "RAN resources exist but the Core has no logical session"),
            ("Access resource mismatch", "Applied RAN resources disagree with Core expectations"),
            (
                "QoS profile mismatch",
                "Both sides expect active resources but their QoS profiles differ",
            ),
        ],
        [2.0, 4.9],
    )
    p(
        "The verifier does not equate RRC_CONNECTED with PDU_SESSION_ACTIVE. Different labels can describe compatible conditions. A retained session with deactivated access resources is an intentional normal state [4–6]."
    )
    h("How the verifier decides")
    bullets(
        [
            "Correlate snapshots by UE/session and discard duplicate or older node sequence numbers.",
            "Require both node views to be fresh. If a snapshot is older than 60 ms, report insufficient observation evidence rather than desynchronization.",
            "Require overlapping stable intervals: both capture times must follow the latest reported state change. The simulator supplies a common clock.",
            "Allow an operation to remain pending until its deadline. Afterward require a post-deadline Core observation.",
            "Raise one alarm per persistent rule after the configured grace period, 12 ms by default. Clear persistence evidence when views become incompatible in time or consistent.",
        ]
    )
    h("Why the fault model matters")
    p(
        "NGAP signaling uses reliable SCTP transport [7]. Our loss scenarios represent an application request or completion that is not processed. They do not claim that ordinary packet loss silently defeats reliable delivery. Reordering occurs in the monitoring channel. All timeouts and commit semantics are experimental choices, not mandated 3GPP values."
    )

    page("Implementation and experiment setup")
    table(
        ["Component", "Purpose"],
        [
            (
                "Python and SimPy",
                "Execute independently modeled node procedures in virtual time [8]",
            ),
            ("Verifier", "Consume snapshots and issue rule-specific alarms"),
            ("Fault channels", "Separate processing failures from telemetry faults"),
            (
                "Recovery adapter",
                "Query, check versions, reconcile and verify with bounded retries",
            ),
            ("pandas and Matplotlib", "Aggregate actual measurements and plot comparisons"),
            ("Streamlit", "Show observed profiles, events, alarms and downloadable evidence [9]"),
            ("uv Ruff ty pytest", "Lock dependencies, check code and test behavior [10]"),
        ],
        [1.8, 5.1],
    )
    h("Recovery procedure")
    p(
        "After an alarm, the custom adapter queries both nodes. It refuses a pending operation, checks that versions have not changed before writing, and aligns RAN state with the Core's committed authorization. A second query verifies convergence. A failed job retries at most twice. This adapter is an experimental hook, not a standardized Context Sync Request."
    )
    h("Default experiment configuration")
    p(
        "Each run has 3 UEs with 2 sessions each, a 540 ms horizon, 2 ms monitor/oracle ticks, 15 ms snapshot polling, a 110 ms procedure timeout and a 12 ms persistence grace. Basic message-leg delay is uniformly 3–9 ms. The delay-only control uses 18–42 ms legs. The load scenario adds a shared three-worker processing queue. Fault eligibility is sampled independently per session with probability 0.7. These are explicit laboratory parameters."
    )
    h("Scenario coverage")
    p(
        "The 12 scenarios are clean operation, benign delay, unprocessed request, unprocessed modification completion, unprocessed setup completion, unprocessed release completion, reordered observations, queued high load, observation outage, access idle/resume, local intra-gNB cell change, and unavailable recovery adapter. High load can create genuine divergence after a timeout. The local mobility control does not implement handover signaling."
    )

    page("Evaluation and detector results")
    h("Ground truth and metric definitions")
    p(
        "The oracle reads actual node states and labels contradiction episodes after pending procedures end. The verifier only sees snapshots. An alarm matches a true episode only when session, rule and time interval agree. An alarm during a valid transient is a false positive, even if a failure appears later. Both oracle and detector use the same semantic rules, so this evaluates the monitor under the model rather than independently proving the model correct."
    )
    p(
        "Precision counts correctly matched alarms. Recall counts detected truth episodes. False-positive rate is the fraction of sessions with no truth episode in a run that nevertheless receive a false alarm. Duplicate alarms are reported separately. Detection delay starts at the oracle's episode start. Undefined denominators are N/A, and aggregate rates pool counts."
    )
    p(
        "The main suite has 1,920 runs: 12 scenarios × 40 seeds × four combinations: immediate, fixed 25 ms, aware without recovery, and aware with recovery. Seeds are paired. The following rows use the aware detector without recovery."
    )
    rows = []
    for label, scenario in (
        ("Setup completion loss", "setup_loss_ack"),
        ("QoS completion loss", "loss_ack"),
        ("Release completion loss", "release_loss_ack"),
        ("Queued high load", "high_load"),
        ("Observation outage", "observation_outage"),
    ):
        r = select(scenario)
        rows.append((label, r["tp"], r["fp"], r["fn"], f"{r['mean_detection_delay_ms']:.2f}"))
    table(["Scenario", "TP", "FP", "FN", "Delay ms"], rows, [3.0, 0.6, 0.6, 0.6, 2.1])
    p(
        "All listed truth episodes were detected in this controlled suite. Clean, benign-delay, request-failure, reordered-observation, idle/resume and local-cell controls had no aware false alarms. Perfect scores on these selected synthetic cases do not estimate real-network accuracy."
    )
    page("Baseline comparison and recovery evidence")
    p(
        "Immediate checking alarms whenever observed profiles differ. The fixed baseline requires 25 ms of mismatch but does not understand pending procedures or align snapshot intervals. Their alarms can precede a genuine fault, so they do not count as detections of a later episode. The aware monitor waits for evidence about a completed operation."
    )
    doc.add_picture(str(ROOT / "results/benchmark/false_positives.png"), width=Inches(6.3))
    p(
        "Figure 1  Immediate and fixed-window monitors mistake valid transitions for faults. High-load rates use only sessions with no truth episode."
    )
    doc.add_picture(str(ROOT / "results/benchmark/recovery.png"), width=Inches(6.3))
    p(
        "Figure 2  Fixed consistency probes improve with timely reconciliation. Observation outages and adapter failures limit the benefit. These are consistency proxies rather than packet-delivery measurements."
    )

    page("Recovery results and timing sensitivity")
    p(
        "Three fixed consistency probes occur at 285, 335 and 510 ms. A probe fails when the actual model has a stable contradiction. These probes do not measure throughput, call drops or user connectivity. Lead time applies only before a failed probe in the same episode. Successful rollback can make a probe pass without re-establishing a session."
    )
    rows = []
    for label, scenario in (
        ("Setup completion loss", "setup_loss_ack"),
        ("QoS completion loss", "loss_ack"),
        ("Release completion loss", "release_loss_ack"),
        ("Queued high load", "high_load"),
        ("Observation outage", "observation_outage"),
        ("Adapter unavailable", "recovery_failure"),
    ):
        before, after = select(scenario), select(scenario, True)
        rows.append(
            (
                label,
                before["failed_probes"],
                after["failed_probes"],
                f"{after['recovery_success_rate'] * 100:.1f}%",
            )
        )
    table(
        ["Scenario", "Failed probes off", "Failed probes on", "Recovery success"],
        rows,
        [2.65, 1.35, 1.35, 1.55],
    )
    p(
        "Each row contains 720 probes, from 40 runs × 6 sessions × 3 probe times. Recovery success counts verified write-and-converge jobs divided by recovery jobs. Jobs that find the context already consistent are not counted as successful writes. The outage delays action until another procedure is near, causing version/pending refusals. An unavailable adapter yields no successful recovery."
    )
    h("Held out sensitivity experiments")
    p(
        "A further 720 runs use seeds 101–120, three scenarios, poll intervals of 5, 15, 30 and 90 ms, and grace periods of 0, 12 and 24 ms. Event-triggered snapshots remain enabled, so the poll interval alone does not equal all observation latency. For QoS completion loss with a 12 ms grace:"
    )
    rows = []
    for r in sensitivity:
        if r["scenario"] == "loss_ack" and r["grace_ms"] == "12":
            rows.append(
                (r["poll_ms"], r["tp"], r["fn"], f"{float(r['mean_detection_delay_ms']):.2f}")
            )
    table(
        ["Poll interval ms", "Detected episodes", "Missed episodes", "Mean delay ms"],
        rows,
        [1.65, 1.75, 1.75, 1.75],
    )
    p(
        "Detection remains model-dependent. Freshness suppression avoids confident alarms during an observation blackout, but it cannot create evidence about an unobserved state."
    )

    page("Limitations and validity of the claims")
    bullets(
        [
            "Abstract protocol model. Setup, modification and release are simplified resource transactions. Full NAS, NGAP, N4, QoS-flow/DRB mappings and rollback behavior are not implemented or certified.",
            "Synthetic observations. The verifier assumes complete internal snapshots with reliable identities, versions and change timestamps. Real packet captures alone may not reveal committed node state.",
            "Clock assumption. Temporal alignment uses a common simulator clock. Production deployment requires clock synchronization and conservative skew bounds.",
            "Shared oracle predicates. The oracle reads actual state independently of telemetry, but uses the detector's semantic contradiction definitions. Standards review and an independent testbed oracle are still needed.",
            "Limited workload diversity. The measured suite has six sessions per run, finite horizons, selected delays and one aggregate processing queue. It does not establish large-network scalability or workload generalization.",
            "Recovery is custom. Core authority is chosen for this model. Different fields in an actual system have different owners. Directly copying context is not a general production recovery method.",
            "Consistency is distinct from availability. Recovery may remove a RAN resource instead of restoring a desired session. A clean failed request is intentionally not a desync. Session success and user demand need separate metrics.",
            "Mobility is a control. The intra-gNB example changes a local cell only. Inter-gNB handover, path switching and AMF relocation are not evaluated.",
            "No learned baseline reproduction. Immediate and fixed-timeout checks are ablations of this task. They are not implementations of the cited LSTM or KPI framework, so no superiority claim against those papers is supported.",
            "Literature access. Publisher summaries and accessible indexed material support the reference comparison. Full-text limitation analysis and a broader novelty review remain necessary for a publication claim.",
        ]
    )
    h("Verified implementation status")
    percentage = coverage["totals"]["percent_covered"]
    p(
        f"The implementation passes 48 behavioral and integration tests, including the Streamlit dashboard. Line coverage for the core package is {percentage:.1f}%. Ruff and ty checks pass. The exported main and sensitivity experiments contain 2,640 runs. Local wall time is a software-runtime measurement, not a production detection benchmark. The included GitHub workflow has not been executed on GitHub."
    )
    h("What would make the next validation stronger")
    p(
        "Instrument a laboratory RAN and Core, map actual identifiers and field ownership, replay measured procedures through the verifier, and evaluate session success separately from context convergence. Add independent protocol checks and test clock skew, missing telemetry, concurrent operations and genuine handover before attempting production recovery."
    )

    page("Viva answers and demonstration guide")
    questions = [
        (
            "What is the project?",
            "A per-session consistency monitor that distinguishes valid protocol transitions from persistent RAN–Core context disagreement, then evaluates reconciliation in simulation.",
        ),
        (
            "What is the base reference?",
            "We selected the ICIN 2020 data-consistency paper as the conceptual basis. The 2023 and 2025 anomaly papers support the monitoring comparison. A full-paper replication is not claimed.",
        ),
        (
            "What is our novelty?",
            "An executable integration of per-session semantic rules, procedure-aware temporal evidence, explainable alarms and bounded recovery experiments. Worldwide novelty is not yet established.",
        ),
        (
            "Why use rules instead of ML?",
            "Our target is a specified contradiction. Rules expose the cause and avoid a training-data requirement. We have not shown that rules outperform ML on general anomalies.",
        ),
        (
            "Why does a delayed message not trigger an alarm?",
            "Known operations can remain pending until their deadline. The verifier also requires fresh, aligned observations and persistent mismatch evidence.",
        ),
        (
            "How do we know which faults are real?",
            "Actual node states define post-procedure truth episodes. Fault tags never enter the detector. Both use the same modeled semantics, which is a stated validation limitation.",
        ),
        (
            "Is recovery standardized?",
            "No. The custom simulator adapter queries, version-checks, reconciles and verifies. A real deployment needs field ownership and supported recovery procedures.",
        ),
        (
            "What do the results prove?",
            "The design works on the declared synthetic scenarios and reduces mismatch exposure when the adapter is available. It does not prove uninterrupted connectivity or operator-network accuracy.",
        ),
    ]
    for question, answer in questions:
        h(question.replace("?", ""))
        p(answer)
    h("Run and explain the demo")
    p(
        "From the project folder, run uv sync --locked --all-groups, then uv run streamlit run app.py. Select loss_ack with the aware detector and enable recovery. Explain the pending operation, timeout, aligned observations, alarm evidence and recovery outcome. Repeat with recovery disabled, then show idle_resume and recovery_failure. The README contains commands for every exported experiment."
    )

    page("References")
    p(
        "Reference versions and links are fixed for the project explanation. Bracketed numbers in this document identify the sources below. The implementation's timeouts, QoS integers, commit rules and custom recovery policy are laboratory abstractions."
    )
    refs = [
        (
            "[1] J. Sid-Otmane, S. Imadali, F. Martelli and M. Shapiro. Data Consistency in the 5G Specification. ICIN 2020, pp. 110–117. DOI 10.1109/ICIN48450.2020.9059408. ",
            "Publisher record",
            "https://ieeexplore.ieee.org/document/9059408/",
        ),
        (
            "[2] Anomaly Detection using a Semi-Supervised Deep Learning Model on Open 5G Core Metrics during User-Equipment Registration. IEEE CSCN 2023. DOI 10.1109/CSCN60443.2023.10453203. ",
            "Publisher record",
            "https://ieeexplore.ieee.org/document/10453203/",
        ),
        (
            "[3] A Framework for AI-Based Large-Scale Time-Series Anomaly Detection in 5G Core Networks. IEEE ICTC 2025. DOI 10.1109/ICTC66702.2025.11388910. IEEE record added in February 2026. ",
            "Publisher record",
            "https://ieeexplore.ieee.org/abstract/document/11388910",
        ),
        (
            "[4] 3GPP TS 23.502 V16.18.0. Procedures for the 5G System. Sections 4.3.2, 4.3.3 and 4.3.4 cover session establishment, modification and release. ",
            "ETSI specification",
            "https://www.etsi.org/deliver/etsi_ts/123500_123599/123502/16.18.00_60/ts_123502v161800p.pdf",
        ),
        (
            "[5] 3GPP TS 38.413 V16.5.0. NG Application Protocol. Sections 8.2.1, 8.2.2 and 8.2.3 cover resource setup, release and modification. ",
            "ETSI specification",
            "https://www.etsi.org/deliver/etsi_ts/138400_138499/138413/16.05.00_60/ts_138413v160500p.pdf",
        ),
        (
            "[6] 3GPP TS 38.300 V16.2.0. NR and NG-RAN Overall Description. Background for access states and mobility. ",
            "ETSI specification",
            "https://www.etsi.org/deliver/etsi_ts/138300_138399/138300/16.02.00_60/ts_138300v160200p.pdf",
        ),
        (
            "[7] 3GPP TS 38.412 V16.0.0. NG signalling transport. Sections 4.1 and 7 describe reliable NGAP signaling over SCTP. ",
            "ETSI specification",
            "https://www.etsi.org/deliver/etsi_ts/138400_138499/138412/16.00.00_60/ts_138412v160000p.pdf",
        ),
        (
            "[8] SimPy. Discrete-event simulation framework. ",
            "Official documentation",
            "https://simpy.readthedocs.io/en/latest/index.html",
        ),
        (
            "[9] Streamlit. Interactive Python application framework. ",
            "Official documentation",
            "https://docs.streamlit.io/",
        ),
        (
            "[10] Astral uv. Python project and dependency management. ",
            "Official documentation",
            "https://docs.astral.sh/uv/",
        ),
    ]
    for description, label, url in refs:
        para = p(description)
        ref_link(para, label, url)
    p(
        "Experimental evidence: results/benchmark/runs.csv, summary.csv, metadata.json, results/sensitivity.csv and results/test-coverage.json. The metadata records configuration, software environment and source hashes. Results are local measurements from this implementation."
    )
    output = ROOT / "deliverables/Project_Pre_Report.docx"
    output.parent.mkdir(exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == "__main__":
    build()
